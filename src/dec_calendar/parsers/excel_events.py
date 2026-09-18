from __future__ import annotations

import re
from datetime import date, datetime
from pathlib import Path
from typing import Iterable, Optional

from openpyxl import load_workbook

from dec_calendar.models import CalendarEventDraft, EventKind

DATE_RE = re.compile(r"^(\d{2})\.(\d{2})\.(\d{4})$")

# Excel "Тип" → EventKind (+ optional extra labels)
GEORGIA_HINTS = ("грузия",)
HOLIDAY_HINTS = ("грузия",)


def _parse_concrete_date(raw: object) -> Optional[date]:
    if raw is None:
        return None
    if isinstance(raw, datetime):
        return raw.date()
    if isinstance(raw, date):
        return raw
    text = str(raw).strip()
    m = DATE_RE.match(text)
    if not m:
        return None
    day, month, year = map(int, m.groups())
    return date(year, month, day)


def map_excel_type_to_kind(type_raw: object) -> EventKind:
    text = (str(type_raw or "").strip().lower())
    if any(h in text for h in HOLIDAY_HINTS):
        return EventKind.HOLIDAY
    return EventKind.PROMO


def extra_labels_for_excel_type(type_raw: object) -> tuple[str, ...]:
    text = str(type_raw or "").strip().lower()
    labels: list[str] = []
    if "кофе" in text:
        labels.append("coffee")
    if "сезон" in text:
        labels.append("seasonal")
    if "коммерч" in text:
        labels.append("commercial")
    if "dk" in text or "внутрен" in text:
        labels.append("dk-internal")
    if "международ" in text:
        labels.append("international")
    if "грузия" in text:
        labels.append("georgia")
    return tuple(labels)


def holiday_checklist(title: str) -> str:
    return (
        f"Событие: {title}\n\n"
        "Чеклист подготовки акции / ивента:\n"
        "- [ ] Определить формат (акция / контент / ивент в точке)\n"
        "- [ ] Согласовать оффер и креатив\n"
        "- [ ] Подготовить стоки / меню при необходимости\n"
        "- [ ] Запланировать публикации\n"
        "- [ ] Назначить ответственного\n"
    )


def parse_excel_calendar(path: Path, *, sheet_name: Optional[str] = None) -> list[CalendarEventDraft]:
    """Parse dated + needs-date rows from the main calendar sheet."""
    wb = load_workbook(path, data_only=True)
    if sheet_name:
        ws = wb[sheet_name]
    else:
        # First sheet is the main calendar in the DimKava workbook.
        ws = wb[wb.sheetnames[0]]

    drafts: list[CalendarEventDraft] = []
    for i, row in enumerate(ws.iter_rows(values_only=True), start=1):
        if i == 1:
            continue
        if not row or all(c is None or str(c).strip() == "" for c in row[:4]):
            continue
        date_cell, _month, title_cell, type_cell = row[0], row[1], row[2], row[3]
        title = str(title_cell or "").strip()
        if not title:
            continue

        due = _parse_concrete_date(date_cell)
        kind = map_excel_type_to_kind(type_cell)
        extras = extra_labels_for_excel_type(type_cell)
        type_label = str(type_cell or "").strip()

        if due is None:
            drafts.append(
                CalendarEventDraft(
                    summary=f"[Промо] {title}" if kind is EventKind.PROMO else f"[Праздник] {title}",
                    kind=kind,
                    due_date=None,
                    description=(
                        f"Событие: {title}\n"
                        f"Тип (Excel): {type_label}\n"
                        f"Дата в источнике: {date_cell}\n\n"
                        "Дата подвижная / размытая — проставьте duedate в Jira и снимите needs-date."
                    ),
                    extra_labels=extras,
                    source=f"excel:{path.name}:{i}",
                    needs_date=True,
                )
            )
            continue

        prefix = "[Праздник]" if kind is EventKind.HOLIDAY else "[Промо]"
        drafts.append(
            CalendarEventDraft(
                summary=f"{prefix} {title}",
                kind=kind,
                due_date=due,
                description=holiday_checklist(title) + f"\nТип (Excel): {type_label}\nИсточник: {path.name}",
                extra_labels=extras,
                source=f"excel:{path.name}:{i}",
                needs_date=False,
            )
        )
    return drafts


def find_excel_workbook(repo_root: Path) -> Path:
    matches = sorted(repo_root.glob("*.xlsx"))
    if not matches:
        raise FileNotFoundError(f"No .xlsx found in {repo_root}")
    return matches[0]


def merge_unique_by_summary(drafts: Iterable[CalendarEventDraft]) -> list[CalendarEventDraft]:
    """Dedupe by summary: preserve first-seen order, last draft wins on content."""
    by_summary: dict[str, CalendarEventDraft] = {}
    order: list[str] = []
    for d in drafts:
        if d.summary not in by_summary:
            order.append(d.summary)
        by_summary[d.summary] = d
    return [by_summary[s] for s in order]
